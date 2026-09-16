"""Gold: app-ready tables built from the silver Parquet.

    python -m src.gold            # from pipeline/, after python -m src.build

Silver is one long observation table with full provenance. Gold reshapes it into
what the Atlas queries, and nothing else:

- ``organisation``    one row per (org_level, org_code): name, ONS code, parent, the
                      periods it exists for, and the period its series breaks (if any)
- ``measure``         one row per distinct measure/breakdown/dimension combination,
                      with a stable ``measure_key``, a plain-English description from
                      the data dictionary, a unit, and its cross-era comparability
- ``period``          one row per reporting period, with its era and source release
- ``observation``     the fact table: latest-release-wins silver, keyed on
                      (period_end, org_level, org_code, measure_key)
- ``diagnosis_rate``  the headline, wide: register / estimate / rate / CI per
                      organisation-period, at every level the publisher gives it
- ``series_break``    the breaks table, as built by ``series_breaks``
- ``geometry``        Sub-ICB boundaries (April 2026, EPSG:4326) as GeoJSON text,
                      joined to org_code via ONS code; ICB and region outlines are
                      dissolved from these in PostGIS at load time

Everything is written as Parquet under ``data/processed/gold/``; ``load_postgis``
pushes the same frames into PostGIS.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

from .measure_crosswalk import ALL, NOT_APPLICABLE
from .org_crosswalk import ENGLAND_ODS_CODE, ENGLAND_ONS_CODE
from .silver_loader import classify_file, read_raw, read_silver
from .silver_schema import DIMENSION_COLUMNS

PIPELINE_DIR = Path(__file__).resolve().parents[1]
RAW_ROOT = PIPELINE_DIR / "data" / "raw" / "pcdd"
SILVER_DIR = PIPELINE_DIR / "data" / "processed" / "silver"
GOLD_DIR = PIPELINE_DIR / "data" / "processed" / "gold"
GEOJSON = next(iter((PIPELINE_DIR / "data" / "raw").glob("Sub_Integrated_Care_Board_Locations_*.geojson")), None)
BOUNDARY_VERSION = "2026-04"   # April 2026 boundaries, i.e. post-reorganisation

RATE_MEASURES = {
    "DEMENTIA_REGISTER_65_PLUS": "register_65_plus",
    "DEMENTIA_ESTIMATE_65_PLUS": "estimate_65_plus",
    "DIAG_RATE_65_PLUS": "diag_rate",
    "DIAG_RATE_65_PLUS_LL": "diag_rate_ll",
    "DIAG_RATE_65_PLUS_UL": "diag_rate_ul",
}

GOLD_TABLES = ("organisation", "measure", "period", "observation", "diagnosis_rate", "series_break", "geometry")


# --------------------------------------------------------------------------------------
# measure
# --------------------------------------------------------------------------------------

def measure_key(row) -> str:
    """Stable string id: ``MEASURE:BREAKDOWN`` plus any non-ALL dimension as
    ``dim=value``. E.g. ``DEMENTIA_REGISTER:AGE_GENDER:age=65_69:gender=Female``."""
    parts = [row["measure"], row["breakdown"]]
    parts += [f"{d}={row[d]}" for d in DIMENSION_COLUMNS if row[d] != ALL]
    return ":".join(parts)


def dictionary_descriptions(path: Path) -> tuple[dict, dict]:
    """(measure/breakdown -> description, dimension token -> description) from the
    Era-B data dictionary. Tolerates the dictionary's own naming slips (``REVIEW``
    for ``REVIEWS``, ``PAT_LIST_65_PLUS`` as a measure)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    measures, tokens = {}, {}
    for sheet in ("pcdem-sub-icb", "pcdem-practice", "pcdem-nhs-rate"):
        for row in wb[sheet].iter_rows(values_only=True):
            if not row or row[0] not in ("MEASURE/BREAKDOWN", "MEASURE"):
                continue
            if row[0] == "MEASURE":
                measure, breakdown, desc = row[2], NOT_APPLICABLE, row[3]
            else:
                measure, breakdown, desc = row[2], row[3], row[4]
            measure = {"REVIEW": "REVIEWS", "PAT_LIST_65_PLUS": "PAT_LIST"}.get(measure, measure)
            measures[(measure, breakdown)] = str(desc).strip()
    for sheet in ("Age_Gender categories", "Ethnicity categories", "Dementia Type categories", "Residential Type categories"):
        for row in wb[sheet].iter_rows(values_only=True):
            if not row or row[0] in (None, "AGE", "ETHNICITY", "DEMENTIA_TYPE", "RESIDENTIAL_TYPE"):
                continue
            if sheet.startswith("Age"):
                tokens[("age", str(row[0]))] = f"aged {row[2].split('aged ')[-1]}" if row[2] else str(row[0])
            else:
                dim = {"Ethnicity": "ethnicity", "Dementia": "dementia_type", "Residential": "residential_type"}[sheet.split()[0]]
                tokens[(dim, str(row[0]))] = str(row[1]).strip()
    return measures, tokens


def measure_dimension(silver: pd.DataFrame, dictionary_path: Path) -> pd.DataFrame:
    descriptions, tokens = dictionary_descriptions(dictionary_path)
    cols = ["measure", "breakdown", *DIMENSION_COLUMNS, "comparability"]
    dim = silver[cols].drop_duplicates(["measure", "breakdown", *DIMENSION_COLUMNS]).reset_index(drop=True)
    dim["measure_key"] = dim.apply(measure_key, axis=1)

    def describe(row):
        base = descriptions.get((row["measure"], row["breakdown"]))
        if base is None:
            base = descriptions.get((row["measure"], NOT_APPLICABLE), f"{row['measure']} ({row['breakdown']})")
        extras = []
        for d in DIMENSION_COLUMNS:
            if row[d] == ALL:
                continue
            if d == "gender":
                extras.append(row[d].lower())
            else:
                extras.append(tokens.get((d, row[d]), row[d]))
        return base if not extras else f"{base} — {', '.join(extras)}"

    dim["description"] = dim.apply(describe, axis=1)
    dim["unit"] = "count"
    dim.loc[dim["measure"].str.startswith("DIAG_RATE"), "unit"] = "percent"
    dim.loc[dim["measure"] == "DEMENTIA_ESTIMATE_65_PLUS", "unit"] = "estimate"
    dim["is_additive"] = dim["unit"] == "count"
    return dim[["measure_key", "measure", "breakdown", *DIMENSION_COLUMNS, "description", "unit",
                "is_additive", "comparability"]].sort_values("measure_key").reset_index(drop=True)


# --------------------------------------------------------------------------------------
# organisation
# --------------------------------------------------------------------------------------

_NAME_SOURCES = {
    # family -> (org_type column, code column, name column)
    "nhs_rate": ("ORG_TYPE", "ORG_CODE", "NAME"),
    "la_rate": ("ORG_TYPE", "ONS_CODE", "NAME"),
    "sub_icb_consolidated": ("ORG_TYPE", "ODS_CODE", "NAME"),
    "practice_measures": ("ORG_TYPE", "ODS_CODE", "NAME"),
}


def organisation_names(raw_root: Path) -> pd.DataFrame:
    """(org_level, org_code) -> display name, latest release wins. Silver drops names
    on purpose (they change casing between releases and must never be joined on);
    gold puts them back for display."""
    from .org_crosswalk import canonical_ltla_ons_code, to_org_level

    frames = []
    for release_dir in sorted(p for p in raw_root.iterdir() if p.is_dir()):
        for path in release_dir.glob("*.csv"):
            family = classify_file(path.name)
            if family not in _NAME_SOURCES:
                continue
            type_col, code_col, name_col = _NAME_SOURCES[family]
            raw = read_raw(path)[[type_col, code_col, name_col]].drop_duplicates()
            frames.append(pd.DataFrame({
                "org_level": raw[type_col].map(to_org_level),
                "org_code": raw[code_col].map(canonical_ltla_ons_code) if family == "la_rate" else raw[code_col],
                "name": raw[name_col].str.strip().str.title().str.replace(r"\bIcb\b", "ICB", regex=True)
                        .str.replace(r"\bNhs\b", "NHS", regex=True),
                "source_release": release_dir.name,
            }))
    names = pd.concat(frames, ignore_index=True).sort_values("source_release")
    return names.drop_duplicates(["org_level", "org_code"], keep="last")[["org_level", "org_code", "name"]]


def organisation_dimension(silver: pd.DataFrame, hierarchy: pd.DataFrame, mapping: pd.DataFrame,
                           breaks: pd.DataFrame, names: pd.DataFrame) -> pd.DataFrame:
    orgs = (silver.groupby(["org_level", "org_code"], dropna=False)
            .agg(ons_code=("ons_code", "last"), first_period=("period_end", "min"),
                 last_period=("period_end", "max"), latest_release=("source_release", "max"))
            .reset_index())
    orgs["is_current"] = orgs["latest_release"] == silver["source_release"].max()

    # Parents, from the latest snapshot that mentions the organisation.
    latest = hierarchy.sort_values("source_release").drop_duplicates("sub_icb_code", keep="last")
    parent = {}
    parent.update({("sub_icb", r.sub_icb_code): ("icb", r.icb_code) for r in latest.itertuples()})
    parent.update({("icb", r.icb_code): ("nhs_region", r.region_code) for r in latest.itertuples()})
    parent.update({("nhs_region", r.region_code): ("country", ENGLAND_ODS_CODE) for r in latest.itertuples()})
    latest_map = mapping.sort_values("source_release").drop_duplicates("practice_code", keep="last")
    parent.update({("practice", r.practice_code): ("sub_icb", r.sub_icb_code)
                   for r in latest_map.itertuples() if not r.unmapped})
    parent.update({("gor", code): ("country", ENGLAND_ONS_CODE)
                   for code in orgs.loc[orgs["org_level"] == "gor", "org_code"]})
    parents = orgs.apply(lambda r: parent.get((r["org_level"], r["org_code"]), (pd.NA, pd.NA)), axis=1)
    orgs["parent_org_level"] = [p[0] for p in parents]
    orgs["parent_org_code"] = [p[1] for p in parents]

    org_breaks = (breaks[breaks["scope"] == "org"].groupby(["org_level", "org_code"])["effective_from"].min()
                  .rename("series_break_from").reset_index())
    orgs = orgs.merge(org_breaks, on=["org_level", "org_code"], how="left")
    orgs = orgs.merge(names, on=["org_level", "org_code"], how="left")
    orgs.loc[(orgs["org_level"] == "country") & orgs["name"].isna(), "name"] = "England"
    cols = ["org_level", "org_code", "ons_code", "name", "parent_org_level", "parent_org_code",
            "first_period", "last_period", "is_current", "series_break_from"]
    return orgs[cols].sort_values(["org_level", "org_code"]).reset_index(drop=True)


# --------------------------------------------------------------------------------------
# period, observation, diagnosis_rate, geometry
# --------------------------------------------------------------------------------------

def period_dimension(latest: pd.DataFrame) -> pd.DataFrame:
    periods = (latest.groupby("period_end")
               .agg(publication_era=("publication_era", "first"), source_release=("source_release", "max"),
                    dictionary_version=("dictionary_version", "first"), n_rows=("value_raw", "size"))
               .reset_index().sort_values("period_end"))
    periods["period"] = periods["period_end"].dt.strftime("%Y-%m")
    return periods[["period_end", "period", "publication_era", "source_release", "dictionary_version", "n_rows"]].reset_index(drop=True)


def observation_fact(latest: pd.DataFrame, measures: pd.DataFrame) -> pd.DataFrame:
    keyed = latest.merge(measures[["measure_key", "measure", "breakdown", *DIMENSION_COLUMNS]],
                         on=["measure", "breakdown", *DIMENSION_COLUMNS], how="left")
    if keyed["measure_key"].isna().any():
        raise ValueError("observation rows with no measure_key - measure dimension is incomplete")
    out = keyed.rename(columns={"value_num": "value", "value_num_lower": "value_lower", "value_num_upper": "value_upper"})
    cols = ["period_end", "org_level", "org_code", "measure_key", "value", "value_state", "value_lower",
            "value_upper", "is_derived", "dq_flag", "comparability", "source_release"]
    out = out[cols].sort_values(["period_end", "org_level", "org_code", "measure_key"]).reset_index(drop=True)
    if out.duplicated(["period_end", "org_level", "org_code", "measure_key"]).any():
        raise ValueError("duplicate observation keys after latest-release resolution")
    return out


def diagnosis_rate_wide(latest: pd.DataFrame) -> pd.DataFrame:
    rate = latest[latest["measure"].isin(RATE_MEASURES) & ~latest["is_derived"].astype(bool)]
    wide = (rate.pivot_table(index=["period_end", "org_level", "org_code"], columns="measure",
                             values="value_num", aggfunc="first")
            .rename(columns=RATE_MEASURES).reset_index())
    dq = rate.groupby(["period_end", "org_level", "org_code"])["dq_flag"].max().rename("dq_flag").reset_index()
    wide = wide.merge(dq, on=["period_end", "org_level", "org_code"])
    cols = ["period_end", "org_level", "org_code", *RATE_MEASURES.values(), "dq_flag"]
    wide.columns.name = None
    return wide[cols].sort_values(["period_end", "org_level", "org_code"]).reset_index(drop=True)


def geometry_table(geojson_path: Path, organisation: pd.DataFrame) -> pd.DataFrame:
    """Sub-ICB polygons keyed by org_code, via the ONS code the boundary file uses."""
    features = json.load(open(geojson_path, encoding="utf-8"))["features"]
    subs = organisation[organisation["org_level"] == "sub_icb"].set_index("ons_code")["org_code"]
    rows = []
    for f in features:
        ons = f["properties"]["SICBL26CD"]
        rows.append({"org_level": "sub_icb", "org_code": subs.get(ons, pd.NA), "ons_code": ons,
                     "name_in_boundary_file": f["properties"]["SICBL26NM"],
                     "boundary_version": BOUNDARY_VERSION,
                     "geometry_geojson": json.dumps(f["geometry"], separators=(",", ":"))})
    return pd.DataFrame(rows).sort_values("ons_code").reset_index(drop=True)


# --------------------------------------------------------------------------------------
# build
# --------------------------------------------------------------------------------------

def latest_dictionary(raw_root: Path) -> Path:
    candidates = sorted(raw_root.glob("*/PCDD-*-data-dictionary.xlsx"))
    if not candidates:
        raise FileNotFoundError("no PCDD data dictionary under raw releases")
    return candidates[-1]


def build_gold(silver_dir: Path = SILVER_DIR, raw_root: Path = RAW_ROOT,
               geojson_path: Path | None = GEOJSON) -> dict[str, pd.DataFrame]:
    silver = read_silver(silver_dir / "pcdd_observations.parquet")
    latest = read_silver(silver_dir / "pcdd_latest.parquet")
    mapping = pd.read_parquet(silver_dir / "pcdd_mapping.parquet")
    hierarchy = pd.read_parquet(silver_dir / "pcdd_hierarchy.parquet")
    breaks = pd.read_parquet(silver_dir / "pcdd_series_breaks.parquet")

    measures = measure_dimension(silver, latest_dictionary(raw_root))
    organisation = organisation_dimension(silver, hierarchy, mapping, breaks, organisation_names(raw_root))
    gold = {
        "organisation": organisation,
        "measure": measures,
        "period": period_dimension(latest),
        "observation": observation_fact(latest, measures),
        "diagnosis_rate": diagnosis_rate_wide(latest),
        "series_break": breaks,
        "geometry": geometry_table(geojson_path, organisation) if geojson_path else pd.DataFrame(),
    }
    return gold


def write_gold(gold: dict[str, pd.DataFrame], out_dir: Path = GOLD_DIR) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, frame in gold.items():
        frame.to_parquet(out_dir / f"{name}.parquet", index=False)


def main(argv: list[str] | None = None) -> int:
    gold = build_gold()
    write_gold(gold)
    for name in GOLD_TABLES:
        print(f"{name:15s} {len(gold[name]):>9,} rows")
    print(f"written to {GOLD_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
